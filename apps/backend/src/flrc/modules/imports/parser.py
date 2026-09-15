from __future__ import annotations

import hashlib
import io
import re
import zipfile
from itertools import islice

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field, ValidationError

from flrc.modules.academics.class_names import (
    MAX_GRADE,
    MIN_GRADE,
    PREP_GRADE,
    SECTION_MAX_LENGTH,
    normalize_section,
)
from flrc.modules.administration.names import search_key

HEADER_ALIASES = {
    "school_number": {"okul no", "öğrenci no", "ogrenci no", "numara", "no"},
    "full_name": {"ad soyad", "adı soyadı", "öğrenci adı soyadı", "ogrenci adi soyadi"},
    "class_combined": {"sınıf/şube", "sinif/sube", "sınıf şube", "class"},
    "grade_level": {"sınıf", "sinif"},
    "section": {"şube", "sube"},
    "language": {"yabancı dil", "2. yabancı dil", "ikinci yabancı dil", "l2"},
}
FOOTER_WORDS = ("toplam", "sınıf mevcudu", "sinif mevcudu", "öğrenci sayısı")
CLASS_RE = re.compile(r"^\s*([1-8])\s*[-/]?\s*([A-Za-zÇĞİÖŞÜçğıöşü])\s*$")
# Hazırlık classes have a name, not a number: "Hazırlık Bulut", "Hazırlık/Bulut"
# or just "Bulut" (ADR-065). The prefix tolerates any Turkish or ASCII i.
PREP_PREFIX_RE = re.compile(r"^haz[ıiİI]rl[ıiİI]k(?=$|[\s/-])", re.IGNORECASE)
PREP_NAME_RE = re.compile(r"^[^\W\d_]+(?: [^\W\d_]+)*$")
MAX_ARCHIVE_ENTRIES = 1_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ARCHIVE_ENTRY_BYTES = 20 * 1024 * 1024
MAX_COMPRESSION_RATIO = 200
MAX_WORKSHEETS = 32
MAX_WORKSHEET_ROWS = 10_000
MAX_WORKSHEET_COLUMNS = 100
HEADER_SCAN_ROWS = 25
REQUIRED_XLSX_MEMBERS = {"[Content_Types].xml", "xl/workbook.xml"}


class InvalidWorkbook(ValueError):
    pass


def validate_xlsx_container(data: bytes) -> None:
    """Reject malformed and excessively expanding XLSX archives before XML parsing."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            names = {entry.filename for entry in entries}
            if not REQUIRED_XLSX_MEMBERS.issubset(names):
                raise InvalidWorkbook("bad_xlsx_container")
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                raise InvalidWorkbook("xlsx_archive_too_complex")
            total_size = sum(entry.file_size for entry in entries)
            total_compressed = sum(entry.compress_size for entry in entries)
            if total_size > MAX_ARCHIVE_UNCOMPRESSED_BYTES or any(
                entry.file_size > MAX_ARCHIVE_ENTRY_BYTES for entry in entries
            ):
                raise InvalidWorkbook("xlsx_archive_too_large")
            if total_size and (
                total_compressed == 0 or total_size / total_compressed > MAX_COMPRESSION_RATIO
            ):
                raise InvalidWorkbook("xlsx_suspicious_compression")
            if any(entry.flag_bits & 0x1 for entry in entries):
                raise InvalidWorkbook("xlsx_encrypted_entry")
    except zipfile.BadZipFile as exc:
        raise InvalidWorkbook("bad_xlsx_container") from exc


class RowModel(BaseModel):
    school_number: int = Field(gt=0)
    full_name: str = Field(min_length=2, max_length=160)
    grade_level: int = Field(ge=MIN_GRADE, le=MAX_GRADE)
    section: str = Field(min_length=1, max_length=SECTION_MAX_LENGTH)
    language: str | None = None
    language_present: bool = False
    sheet: str
    row_number: int


class ImportIssue(BaseModel):
    level: str
    code: str
    sheet: str
    cell: str | None = None
    message: str


class ImportPlan(BaseModel):
    sha256: str
    rows: list[RowModel]
    issues: list[ImportIssue]

    @property
    def has_errors(self) -> bool:
        return any(issue.level == "error" for issue in self.issues)


def norm(value: object) -> str:
    return search_key(str(value or "")).replace("\n", " ")


def header_map(values: list[object]) -> dict[str, int]:
    found: dict[str, int] = {}
    for index, raw in enumerate(values):
        value = norm(raw)
        for canonical, aliases in HEADER_ALIASES.items():
            if value in aliases:
                found[canonical] = index
    return found


def parse_class(value: object, *, bare_prep_name: bool = True) -> tuple[int, str] | None:
    """Read a class from a cell or worksheet title.

    A bare name such as "Bulut" is a Hazırlık class only where the school
    wrote a class, in the class column; a worksheet title is free text and
    must carry the "Hazırlık" prefix so "Students" never becomes a class.
    """
    text = " ".join(str(value or "").split())
    match = CLASS_RE.match(text)
    if match:
        return int(match.group(1)), normalize_section(int(match.group(1)), match.group(2))
    prefix = PREP_PREFIX_RE.match(text)
    if prefix:
        text = text[prefix.end() :]
    elif not bare_prep_name:
        return None
    text = text.strip(" /-")
    if len(text) < 2 or len(text) > SECTION_MAX_LENGTH or not PREP_NAME_RE.match(text):
        return None
    return PREP_GRADE, normalize_section(PREP_GRADE, text)


def parse_language(value: object) -> str | None:
    key = norm(value)
    if not key:
        return None
    if key in {"almanca", "de", "german", "deutsch"}:
        return "german"
    if key in {"fransızca", "fransizca", "fr", "french", "français"}:
        return "french"
    raise ValueError("unknown_language")


def _parse_workbook(data: bytes) -> ImportPlan:
    digest = hashlib.sha256(data).hexdigest()
    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    rows: list[RowModel] = []
    issues: list[ImportIssue] = []

    if len(workbook.worksheets) > MAX_WORKSHEETS:
        raise InvalidWorkbook("xlsx_too_many_worksheets")
    for sheet in workbook.worksheets:
        # A read-only worksheet otherwise trusts the <dimension> the file
        # declares, which is both attacker-controlled and wrong in some real
        # exports: an under-reported dimension would silently import half a
        # roster. Reading the real extent instead means the row and column
        # caps below are counted off the data as it streams past, so they
        # cannot be talked out of by the file.
        sheet.reset_dimensions()
        default_class = parse_class(sheet.title, bare_prep_name=False)
        header_row: int | None = None
        mapping: dict[str, int] = {}
        for row_number, row in enumerate(
            islice(sheet.iter_rows(values_only=True), HEADER_SCAN_ROWS), 1
        ):
            candidate = header_map(list(row))
            if "school_number" in candidate and "full_name" in candidate:
                header_row, mapping = row_number, candidate
                break
        if header_row is None:
            issues.append(
                ImportIssue(
                    level="warning",
                    code="sheet_without_headers",
                    sheet=sheet.title,
                    message="No recognizable student table found; sheet skipped.",
                )
            )
            continue

        for row_number, cells in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1
        ):
            if row_number > MAX_WORKSHEET_ROWS or len(cells) > MAX_WORKSHEET_COLUMNS:
                raise InvalidWorkbook("xlsx_dimensions_too_large")
            values = list(cells)
            joined = " ".join(norm(value) for value in values if value is not None)
            if not joined or any(word in joined for word in FOOTER_WORDS):
                continue

            def at(
                name: str,
                current_mapping: dict[str, int] = mapping,
                current_values: list[object] = values,
            ) -> object:
                index = current_mapping.get(name)
                return (
                    current_values[index]
                    if index is not None and index < len(current_values)
                    else None
                )

            number_raw = at("school_number")
            name_raw = at("full_name")
            if number_raw in (None, "") and name_raw in (None, ""):
                continue
            try:
                numeric = float(str(number_raw).replace(",", "."))
                number = int(numeric)
                if numeric != number:
                    raise ValueError
            except (TypeError, ValueError):
                index = mapping["school_number"] + 1
                issues.append(
                    ImportIssue(
                        level="error",
                        code="bad_school_number",
                        sheet=sheet.title,
                        cell=f"{get_column_letter(index)}{row_number}",
                        message=f"School number is not an integer: {number_raw!r}",
                    )
                )
                continue

            parsed_class = default_class
            if "class_combined" in mapping:
                parsed_class = parse_class(at("class_combined"))
            elif "grade_level" in mapping and "section" in mapping:
                parsed_class = parse_class(f"{at('grade_level')}/{at('section')}")
            if parsed_class is None:
                index = mapping.get("class_combined", mapping.get("grade_level", 0)) + 1
                issues.append(
                    ImportIssue(
                        level="error",
                        code="bad_class",
                        sheet=sheet.title,
                        cell=f"{get_column_letter(index)}{row_number}",
                        message="Could not determine grade and section.",
                    )
                )
                continue
            try:
                rows.append(
                    RowModel(
                        school_number=number,
                        full_name=str(name_raw or "").strip(),
                        grade_level=parsed_class[0],
                        section=parsed_class[1],
                        language=(
                            parse_language(at("language")) if "language" in mapping else None
                        ),
                        language_present="language" in mapping,
                        sheet=sheet.title,
                        row_number=row_number,
                    )
                )
            except (ValidationError, ValueError) as exc:
                issues.append(
                    ImportIssue(
                        level="error",
                        code="row_validation",
                        sheet=sheet.title,
                        cell=f"A{row_number}",
                        message=str(exc),
                    )
                )

    deduped: dict[int, RowModel] = {}
    for row in rows:
        previous = deduped.get(row.school_number)
        if previous is None:
            deduped[row.school_number] = row
            continue
        identical = (
            search_key(previous.full_name) == search_key(row.full_name)
            and previous.grade_level == row.grade_level
            and previous.section == row.section
            and previous.language == row.language
        )
        issues.append(
            ImportIssue(
                level="warning" if identical else "error",
                code="duplicate_identical" if identical else "duplicate_conflict",
                sheet=row.sheet,
                cell=f"A{row.row_number}",
                message=f"School number {row.school_number} appears more than once.",
            )
        )
    return ImportPlan(sha256=digest, rows=list(deduped.values()), issues=issues)


def parse_workbook(data: bytes) -> ImportPlan:
    validate_xlsx_container(data)
    try:
        return _parse_workbook(data)
    except InvalidWorkbook:
        raise
    except Exception as exc:
        # openpyxl/lxml expose several parser-specific exceptions. None are useful
        # to an API caller, and malformed workbook internals must never become a 500.
        raise InvalidWorkbook("malformed_xlsx") from exc
