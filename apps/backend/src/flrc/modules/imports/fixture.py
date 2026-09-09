from pathlib import Path

from openpyxl import Workbook


def make_hostile_fixture(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    # Excel forbids slashes in worksheet titles. The importer accepts this
    # legacy hyphen form and normalizes it to the canonical "5/A" display.
    sheet.title = "5-A"
    sheet.merge_cells("A1:F1")
    sheet["A1"] = "2026-2027 SYNTHETIC STUDENT LIST"
    sheet.append([])
    sheet.append(["Sıra", "Okul No", "Ad Soyad", "Cinsiyet", "Yabancı Dil"])
    sheet.append([1, 51001, "İpek Işık", "K", "Almanca"])
    sheet.append([2, 51002, "Çağrı Şen", "E", "Fransızca"])
    sheet.append([3, 51003, "Öykü Akın", "K", None])
    sheet.append([])
    sheet.append([None, None, "Toplam 3 öğrenci"])

    mixed = workbook.create_sheet("Karma Liste")
    mixed.append(["Kurum raporu — sentetik dışa aktarım"])
    mixed.append([])
    mixed.append(["Öğrenci No", "Öğrenci Adı Soyadı", "Sınıf/Şube", "2. Yabancı Dil"])
    mixed.append([61001, "Yağız Efe", "6/B", "DE"])
    mixed.append([61002, "Gökçe Arı", "6-B", "FR"])

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
