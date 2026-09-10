"""Bulk operations must stay bounded without changing identities or final values."""

from contextlib import contextmanager
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import Session

from flrc.db import models as m
from flrc.modules.administration.names import search_key
from flrc.modules.imports.parser import parse_workbook
from flrc.workers.tasks.exports import build_year_workbook
from tests.conftest import TEST_URL_SYNC, TestSession, cell, save_grid, test_engine


@contextmanager
def count_reads():
    statements = []

    def record(_connection, _cursor, statement, _parameters, _context, _executemany):
        if statement.lstrip().upper().startswith("SELECT"):
            statements.append(statement)

    event.listen(test_engine.sync_engine, "before_cursor_execute", record)
    try:
        yield statements
    finally:
        event.remove(test_engine.sync_engine, "before_cursor_execute", record)


async def test_assignment_matrix_reads_do_not_grow_per_slot(api, world, record_property):
    async with TestSession() as db:
        classes = [
            m.SchoolClass(year_id=world.year, grade_level=5, section=f"P{i}") for i in range(20)
        ]
        db.add_all(classes)
        await db.commit()
    changes = [
        {"class_id": cls.id, "role": role, "user_id": user.id}
        for cls in classes
        for role, user in [("main", world.main), ("skills", world.skills)]
    ]
    with count_reads() as reads:
        async with api(world.admin) as client:
            response = await client.put(
                f"/api/admin/years/{world.year}/assignments", json={"changes": changes}
            )
    assert response.status_code == 200
    slots = {(item["class_id"], item["role"]): item["user_id"] for item in response.json()}
    assert len(slots) == 42  # Includes the existing class, untouched by this request.
    for change in changes:
        assert slots[(change["class_id"], change["role"])] == change["user_id"]
    record_property("select_count", len(reads))
    assert len(reads) <= 5, f"Assignment update issued {len(reads)} SELECTs"


async def test_repeated_assignment_slots_keep_last_change_and_validate_every_change(api, world):
    slot = {"class_id": world.cls, "role": "main"}
    async with api(world.admin) as client:
        response = await client.put(
            f"/api/admin/years/{world.year}/assignments",
            json={
                "changes": [
                    {**slot, "user_id": user_id}
                    for user_id in [None, world.skills.id, world.main.id]
                ]
            },
        )
        rejected = await client.put(
            f"/api/admin/years/{world.year}/assignments",
            json={
                "changes": [
                    {"class_id": world.cls, "role": "german", "user_id": world.main.id},
                    {"class_id": world.cls, "role": "german", "user_id": None},
                ]
            },
        )
        final = await client.get(f"/api/admin/years/{world.year}/assignments")
    assert response.status_code == 200
    assert rejected.status_code == 422
    assert rejected.json()["detail"]["code"] == "teacher_field_mismatch"
    assert final.json() == response.json()


def import_file(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "5-A"
    sheet.append(["Okul No", "Ad Soyad", "Yabancı Dil"])
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    data = buffer.getvalue()
    return {"file": ("synthetic.xlsx", data)}, parse_workbook(data).sha256


async def test_import_existing_roster_uses_bounded_reads_and_preserves_student_ids(
    api, world, record_property
):
    async with TestSession() as db:
        students = [
            m.Student(full_name=f"Synthetic Import {i}", search_name=f"synthetic import {i}")
            for i in range(40)
        ]
        db.add_all(students)
        await db.flush()
        db.add_all(
            [
                m.Enrollment(
                    student_id=student.id,
                    year_id=world.year,
                    class_id=world.cls,
                    school_number=70000 + i,
                )
                for i, student in enumerate(students)
            ]
        )
        db.add_all(
            [
                m.StudentLanguage(student_id=student.id, year_id=world.year, language="german")
                for student in students
            ]
        )
        await db.commit()
        expected_ids = {70000 + i: student.id for i, student in enumerate(students)}
    files, digest = import_file(
        [[70000 + i, f"Renamed Synthetic {i}", "Fransızca"] for i in range(40)]
    )
    with count_reads() as reads:
        async with api(world.admin) as client:
            response = await client.post(
                "/api/admin/import/commit",
                params={"year_id": world.year, "expected_sha256": digest},
                files=files,
            )
    assert response.status_code == 200
    async with TestSession() as db:
        rows = (
            await db.execute(
                select(m.Enrollment, m.Student, m.StudentLanguage)
                .join(m.Student, m.Student.id == m.Enrollment.student_id)
                .join(m.StudentLanguage, m.StudentLanguage.student_id == m.Student.id)
                .where(m.Enrollment.year_id == world.year, m.StudentLanguage.year_id == world.year)
            )
        ).all()
    assert len(rows) == 40
    for enrollment, student, language in rows:
        assert student.id == expected_ids[enrollment.school_number]
        assert student.full_name.startswith("Renamed Synthetic")
        assert language.language == "french"
    record_property("select_count", len(reads))
    assert len(reads) <= 9, f"Import issued {len(reads)} SELECTs"


@pytest.mark.parametrize("pending_count", [1, 2])
async def test_import_claims_numberless_students_only_when_unambiguous(api, world, pending_count):
    async with TestSession() as db:
        students = [
            m.Student(
                full_name="Synthetic Same Name", search_name=search_key("Synthetic Same Name")
            )
            for _ in range(pending_count)
        ]
        db.add_all(students)
        await db.flush()
        db.add_all(
            [
                m.Enrollment(student_id=student.id, year_id=world.year, class_id=world.cls)
                for student in students
            ]
        )
        await db.commit()
        original_ids = {student.id for student in students}
    files, digest = import_file(
        [[70001, "Synthetic Same Name", "Almanca"], [70002, "Synthetic Same Name", "Fransızca"]]
    )
    async with api(world.admin) as client:
        response = await client.post(
            "/api/admin/import/commit",
            params={"year_id": world.year, "expected_sha256": digest},
            files=files,
        )
    assert response.status_code == 200
    async with TestSession() as db:
        rows = list(
            await db.scalars(
                select(m.Enrollment)
                .where(
                    m.Enrollment.year_id == world.year,
                    m.Enrollment.school_number.in_([70001, 70002]),
                )
                .order_by(m.Enrollment.school_number)
            )
        )
    assert len(rows) == 2
    assert (rows[0].student_id in original_ids) == (pending_count == 1)
    assert rows[1].student_id not in original_ids


async def test_year_export_streams_rows_and_finishes_progress(api, world):
    saved = await save_grid(
        api, world.main, world.cls, [cell(world.student_one, world.main_column, 0)]
    )
    assert saved.status_code == 200
    async with TestSession() as db:
        students = [
            m.Student(full_name=f"Synthetic Export {i}", search_name=f"synthetic export {i}")
            for i in range(1200)
        ]
        db.add_all(students)
        await db.flush()
        db.add_all(
            [
                m.Enrollment(
                    student_id=student.id,
                    year_id=world.year,
                    class_id=world.cls,
                    school_number=70000 + i,
                )
                for i, student in enumerate(students)
            ]
        )
        await db.commit()
    engine = create_engine(TEST_URL_SYNC)
    streamed = []

    def record(_connection, _cursor, _statement, _parameters, context, _executemany):
        if context.execution_options.get("yield_per"):
            streamed.append(context.execution_options["yield_per"])

    event.listen(engine, "before_cursor_execute", record)
    try:
        with Session(engine) as db:
            job = m.JobRun(
                kind="year_export", payload={"year_id": world.year}, requested_by=world.admin.id
            )
            db.add(job)
            db.commit()
            data = build_year_workbook(db, world.year, job)
            db.refresh(job)
            assert job.progress == 7
        workbook = load_workbook(BytesIO(data), read_only=True)
        assert len(workbook.sheetnames) == 7
        exported_students = list(workbook["Students"].values)[1:]
        assert len(exported_students) == 1202
        assert exported_students[:2] == [
            (world.student_one, 51001, "Synthetic Student One"),
            (world.student_two, 51002, "Synthetic Student Two"),
        ]
        assert exported_students[-1][1:] == (71199, "Synthetic Export 1199")
        assert list(workbook["Grades"].values)[1][0:9] == (
            51001,
            world.main_column,
            "Main score",
            world.semester,
            "english",
            0,
            None,
            None,
            1,
        )
        assert list(workbook["Audit"].values)[1][3:12] == (
            51001,
            world.main_column,
            None,
            None,
            None,
            0,
            None,
            None,
            False,
        )
        workbook.close()
        assert len(streamed) == 7
        assert all(size <= 1000 for size in streamed)
    finally:
        engine.dispose()


async def test_completeness_and_missing_cells_agree_without_loading_comment_bodies(api, world):
    async with TestSession() as db:

        def column(subject, value_type, position, **kwargs):
            return m.ColumnDefinition(
                semester_id=world.semester,
                grade_level=5,
                subject=subject,
                value_type=value_type,
                owner_role="main" if subject == "english" else subject,
                labels={"tr": "Synthetic field"},
                position=position,
                **kwargs,
            )

        note = column("english", "text", 3)
        inactive = column("english", "score", 4, is_active=False)
        german = column("german", "scale3", 1)
        german_note = column("german", "text", 2)
        french = column("french", "scale3", 1)
        db.add_all([note, inactive, german, german_note, french])
        db.add_all(
            [
                m.StudentLanguage(
                    student_id=world.student_one, year_id=world.year, language="german"
                ),
                m.StudentLanguage(
                    student_id=world.student_two, year_id=world.year, language="french"
                ),
            ]
        )
        await db.flush()
        db.add_all(
            [
                m.GradeValue(
                    student_id=world.student_one,
                    column_definition_id=world.main_column,
                    score=0,
                    updated_by=world.main.id,
                ),
                m.GradeValue(
                    student_id=world.student_one,
                    column_definition_id=world.skills_column,
                    updated_by=world.main.id,
                ),
                m.GradeValue(
                    student_id=world.student_two,
                    column_definition_id=note.id,
                    text_value="",
                    updated_by=world.main.id,
                ),
                m.GradeValue(
                    student_id=world.student_one,
                    column_definition_id=inactive.id,
                    score=99,
                    updated_by=world.main.id,
                ),
                m.GradeValue(
                    student_id=world.student_one,
                    column_definition_id=german.id,
                    scale=1,
                    updated_by=world.main.id,
                ),
                # A grade outside this student's current language roster is ignored.
                m.GradeValue(
                    student_id=world.student_one,
                    column_definition_id=french.id,
                    scale=3,
                    updated_by=world.main.id,
                ),
            ]
        )
        await db.commit()
    with count_reads() as reads:
        async with api(world.admin) as client:
            response = await client.get(
                "/api/coordinator/completeness", params={"semester_id": world.semester}
            )
            assert response.status_code == 200
            rows = {row["subject"]: row for row in response.json()}
            for subject, expected, filled in [
                ("english", 6, 2),
                ("german", 2, 1),
                ("french", 1, 0),
            ]:
                assert rows[subject]["expected_cells"] == expected
                assert rows[subject]["filled_cells"] == filled
                missing = await client.get(
                    f"/api/coordinator/classes/{world.cls}/missing",
                    params={"semester_id": world.semester, "subject": subject},
                )
                assert missing.status_code == 200
                assert len(missing.json()) == rows[subject]["missing_cells"] == expected - filled
                assert all(item["column_id"] != inactive.id for item in missing.json())
    grade_reads = [
        statement.split("FROM")[0] for statement in reads if "FROM grade_values" in statement
    ]
    assert len(grade_reads) == 4
    assert all("text_value" not in statement for statement in grade_reads)
