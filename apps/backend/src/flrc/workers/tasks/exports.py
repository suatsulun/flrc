from io import BytesIO
from typing import Any

from openpyxl import Workbook
from sqlalchemy import Select, select
from sqlalchemy.orm import Session

from flrc.core.spreadsheets import safe_spreadsheet_cell
from flrc.db import models as m
from flrc.db.sync import sync_engine
from flrc.modules.jobs.service import mark_failed, mark_running, mark_succeeded
from flrc.workers.celery import celery_app


def build_year_workbook(db: Session, year_id: int, progress: m.JobRun | None = None) -> bytes:
    year = db.get(m.AcademicYear, year_id)
    if year is None:
        raise ValueError("unknown_year")
    workbook = Workbook(write_only=True)

    def write_sheet(title: str, headers: list[str], statement: Select[Any]) -> None:
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        # Stream both database rows and worksheet XML. Close the cursor before
        # the progress commit, which would invalidate a live server-side cursor.
        with db.execute(statement.execution_options(yield_per=1000)) as rows:
            for row in rows:
                sheet.append([safe_spreadsheet_cell(value) for value in row])
        if progress:
            progress.progress = len(workbook.sheetnames)
            db.commit()

    class_ids = list(db.scalars(select(m.SchoolClass.id).where(m.SchoolClass.year_id == year_id)))
    semester_ids = list(db.scalars(select(m.Semester.id).where(m.Semester.year_id == year_id)))

    students = (
        select(m.Student.id, m.Enrollment.school_number, m.Student.full_name)
        .join(m.Enrollment, m.Enrollment.student_id == m.Student.id)
        .where(m.Enrollment.year_id == year_id)
        .order_by(m.Enrollment.school_number.asc().nulls_last())
    )
    write_sheet("Students", ["id", "school_number", "full_name"], students)

    enrollments = select(
        m.Enrollment.id,
        m.Enrollment.student_id,
        m.Enrollment.year_id,
        m.Enrollment.class_id,
        m.Enrollment.school_number,
    ).where(m.Enrollment.year_id == year_id)
    write_sheet(
        "Enrollments",
        ["id", "student_id", "year_id", "class_id", "school_number"],
        enrollments,
    )

    languages = select(
        m.StudentLanguage.id,
        m.StudentLanguage.student_id,
        m.StudentLanguage.year_id,
        m.StudentLanguage.language,
    ).where(m.StudentLanguage.year_id == year_id)
    write_sheet("Languages", ["id", "student_id", "year_id", "language"], languages)

    assignments = select(
        m.TeachingAssignment.id,
        m.TeachingAssignment.class_id,
        m.TeachingAssignment.role,
        m.TeachingAssignment.user_id,
    ).where(m.TeachingAssignment.class_id.in_(class_ids))
    write_sheet("Assignments", ["id", "class_id", "role", "user_id"], assignments)

    columns = select(
        m.ColumnDefinition.id,
        m.ColumnDefinition.semester_id,
        m.ColumnDefinition.grade_level,
        m.ColumnDefinition.subject,
        m.ColumnDefinition.value_type,
        m.ColumnDefinition.owner_role,
        m.ColumnDefinition.labels["tr"].astext,
        m.ColumnDefinition.position,
        m.ColumnDefinition.is_active,
    ).where(m.ColumnDefinition.semester_id.in_(semester_ids))
    write_sheet(
        "Columns",
        [
            "id",
            "semester_id",
            "grade_level",
            "subject",
            "value_type",
            "owner_role",
            "label_tr",
            "position",
            "is_active",
        ],
        columns,
    )

    grades = (
        select(
            m.Enrollment.school_number,
            m.GradeValue.column_definition_id,
            m.ColumnDefinition.labels["tr"].astext,
            m.ColumnDefinition.semester_id,
            m.ColumnDefinition.subject,
            m.GradeValue.score,
            m.GradeValue.scale,
            m.GradeValue.text_value,
            m.GradeValue.version,
            m.GradeValue.updated_by,
            m.GradeValue.created_at,
            m.GradeValue.updated_at,
        )
        .join(m.Student, m.Student.id == m.GradeValue.student_id)
        .join(m.ColumnDefinition, m.ColumnDefinition.id == m.GradeValue.column_definition_id)
        .join(m.Semester, m.Semester.id == m.ColumnDefinition.semester_id)
        .join(
            m.Enrollment,
            (m.Enrollment.student_id == m.Student.id)
            & (m.Enrollment.year_id == m.Semester.year_id),
        )
        .where(m.ColumnDefinition.semester_id.in_(semester_ids))
    )
    write_sheet(
        "Grades",
        [
            "school_number",
            "column_id",
            "column_label",
            "semester_id",
            "subject",
            "score",
            "scale",
            "text",
            "version",
            "updated_by",
            "created_at",
            "updated_at",
        ],
        grades,
    )

    audit = (
        select(
            m.AuditEntry.id,
            m.AuditEntry.batch_id,
            m.AuditEntry.actor_id,
            m.Enrollment.school_number,
            m.AuditEntry.column_definition_id,
            m.AuditEntry.old_score,
            m.AuditEntry.old_scale,
            m.AuditEntry.old_text,
            m.AuditEntry.new_score,
            m.AuditEntry.new_scale,
            m.AuditEntry.new_text,
            m.AuditEntry.forced,
            m.AuditEntry.created_at,
        )
        .join(m.Student, m.Student.id == m.AuditEntry.student_id)
        .join(m.SaveBatch, m.SaveBatch.id == m.AuditEntry.batch_id)
        .join(
            m.Enrollment,
            (m.Enrollment.student_id == m.Student.id) & (m.Enrollment.year_id == year_id),
        )
        .where(m.SaveBatch.class_id.in_(class_ids))
    )
    write_sheet(
        "Audit",
        [
            "id",
            "batch_id",
            "actor_id",
            "school_number",
            "column_id",
            "old_score",
            "old_scale",
            "old_text",
            "new_score",
            "new_scale",
            "new_text",
            "forced",
            "created_at",
        ],
        audit,
    )
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@celery_app.task(name="exports.year")
def export_year_job(job_id: int) -> None:
    with Session(sync_engine()) as db:
        job = db.get(m.JobRun, job_id)
        if job is None or job.status == "succeeded":
            return
        try:
            year_id = int(str(job.payload["year_id"]))
            mark_running(db, job, total=7)
            data = build_year_workbook(db, year_id, job)
            mark_succeeded(
                db,
                job,
                data=data,
                filename=f"flrc-year-{year_id}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        except Exception:
            db.rollback()
            fresh = db.get(m.JobRun, job_id)
            if fresh is not None:
                mark_failed(db, fresh, "year_export_failed")
            raise
